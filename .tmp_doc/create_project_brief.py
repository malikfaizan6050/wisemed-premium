from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

OUT = Path("docs/WiseMedBilling_Project_Briefing_and_Meeting_QA.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
DARK_BLUE = RGBColor(31, 77, 120)
SLATE = RGBColor(71, 85, 105)
LIGHT_BLUE = "EFF6FF"
LIGHT_GRAY = "F2F4F7"
WHITE = RGBColor(255, 255, 255)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(0.85)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

def font(run, size=11, color=NAVY, bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    return run

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = NAVY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, color, before, after in [
    ("Title", 30, NAVY, 0, 8),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    s = styles[style_name]
    s.font.name = "Calibri"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.font.size = Pt(size)
    s.font.color.rgb = color
    s.font.bold = style_name != "Title"
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for style_name in ["List Bullet", "List Number"]:
    s = styles[style_name]
    s.font.name = "Calibri"
    s.font.size = Pt(11)
    s.font.color.rgb = NAVY
    s.paragraph_format.left_indent = Inches(0.5)
    s.paragraph_format.first_line_indent = Inches(-0.25)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.line_spacing = 1.167

if "Question" not in styles:
    qstyle = styles.add_style("Question", WD_STYLE_TYPE.PARAGRAPH)
    qstyle.font.name = "Calibri"
    qstyle.font.size = Pt(11)
    qstyle.font.bold = True
    qstyle.font.color.rgb = DARK_BLUE
    qstyle.paragraph_format.space_before = Pt(8)
    qstyle.paragraph_format.space_after = Pt(2)
    qstyle.paragraph_format.keep_with_next = True

if "Answer" not in styles:
    astyle = styles.add_style("Answer", WD_STYLE_TYPE.PARAGRAPH)
    astyle.font.name = "Calibri"
    astyle.font.size = Pt(10.5)
    astyle.font.color.rgb = NAVY
    astyle.paragraph_format.left_indent = Inches(0.18)
    astyle.paragraph_format.space_after = Pt(7)
    astyle.paragraph_format.line_spacing = 1.08

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    tblPr.append(tblInd)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tcW.set(qn("w:w"), str(widths[idx]))
            tcW.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Page ")
    font(r, 9, SLATE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
font(header.add_run("WiseMedBilling | Project Briefing"), 9, SLATE, bold=True)
add_page_number(section.footer.paragraphs[0])

def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.75)
    p.add_run(text)
    return p

def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p

def add_callout(label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(f"{label}: "), 10.5, BLUE, bold=True)
    font(p.add_run(text), 10.5, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_qa(question, answer):
    q = doc.add_paragraph(style="Question")
    q.add_run(f"Q. {question}")
    a = doc.add_paragraph(style="Answer")
    a.add_run(f"A. {answer}")

def page_break():
    doc.add_page_break()

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(54)
kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(kicker.add_run("EXECUTIVE PROJECT BRIEFING"), 10, BLUE, bold=True)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("WiseMedBilling")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
font(subtitle.add_run("Complete Platform Overview and CEO/Team Meeting Q&A"), 16, DARK_BLUE, bold=True)
desc = doc.add_paragraph()
desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
desc.paragraph_format.space_after = Pt(58)
font(desc.add_run("Public healthcare revenue-cycle website, lead intake, CRM operations, security controls, integrations, and growth roadmap"), 11, SLATE, italic=True)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(meta.add_run("Prepared for leadership and cross-functional review\nRepository assessment as of August 27, 2026"), 10.5, SLATE)
doc.add_paragraph().paragraph_format.space_after = Pt(42)
add_callout("Purpose", "Provide a meeting-ready, plain-language explanation of what the platform does today, how it is built, what safeguards are present, and what questions leadership should be prepared to answer.")

page_break()
doc.add_heading("Executive summary", level=1)
doc.add_paragraph("WiseMedBilling is a combined public marketing website and internal healthcare revenue-cycle management (RCM) lead operations platform. The public experience explains medical billing and RCM services, captures consultation interest, offers WhatsApp contact options, and supports organic discovery through technical SEO. The internal CRM converts those inquiries into controlled workflows for assignment, follow-up, activity tracking, team oversight, and performance reporting.")
doc.add_paragraph("The solution is implemented as a Next.js 16 application using React 19 and TypeScript. Firebase provides client authentication, server-side identity administration, and Firestore data storage. The application separates public acquisition journeys from authenticated CRM operations through route boundaries, server-side authorization, permissions, and restrictive Firestore rules.")
add_callout("Leadership takeaway", "The project is more than a brochure site: it is an integrated acquisition-to-operations foundation. Before production scale, leadership should confirm the final domain, real support contact details, privacy/compliance review, monitoring, backup/recovery expectations, and ownership of operational processes.")

doc.add_heading("What the platform is designed to achieve", level=2)
for item in [
    "Present WiseMedBilling as a professional US-focused medical billing and healthcare RCM partner.",
    "Educate providers about billing, claims, denials, eligibility, payment posting, analytics, and HIPAA-focused handling.",
    "Convert qualified provider interest into structured leads through consultation and WhatsApp journeys.",
    "Give authorized employees role-appropriate access to leads, notes, activities, assignments, and performance data.",
    "Create a scalable base for operational automation, reporting, specialty expansion, and future integrations."
]: add_bullet(item)

doc.add_heading("Current status at a glance", level=2)
status_rows = [
    ("Public website", "Implemented", "Homepage, About, Services, Solutions, Resources, Contact, Privacy, Terms, and HIPAA pages."),
    ("Lead intake", "Implemented", "Consultation form, reCAPTCHA, validation, scoring, duplicate checks, and CRM creation workflow."),
    ("CRM", "Implemented", "Lead operations, assignments, activities, notifications, roles, users, and analytics views."),
    ("Authentication", "Implemented", "Firebase login, reset request, temporary-password flow, and permission-aware access."),
    ("SEO", "Implemented", "Unique metadata, canonicals, structured data, sitemap, robots rules, and semantic headings."),
    ("Production assurance", "Requires formal sign-off", "Repository validation is not the same as penetration testing, HIPAA certification, legal review, or operational readiness approval.")
]
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
set_table_widths(table, [1900, 1600, 5860])
for idx, text in enumerate(["Area", "Status", "Summary"]):
    set_cell_shading(table.rows[0].cells[idx], LIGHT_GRAY)
    font(table.rows[0].cells[idx].paragraphs[0].add_run(text), 10, NAVY, bold=True)
set_repeat_table_header(table.rows[0])
for area, status, summary in status_rows:
    cells = table.add_row().cells
    for idx, text in enumerate([area, status, summary]):
        font(cells[idx].paragraphs[0].add_run(text), 9.5, NAVY, bold=(idx == 0))
set_table_widths(table, [1900, 1600, 5860])

page_break()
doc.add_heading("1. Public website and customer journey", level=1)
doc.add_paragraph("The public site uses a consistent WiseMedBilling visual system: blue and slate typography, white translucent cards, rounded geometry, soft gradients, clear calls to action, and responsive layouts. Shared Navbar, Footer, call-to-action, legal-page, and public-page components reduce duplication and keep new pages visually aligned.")

doc.add_heading("Public routes", level=2)
routes = [
    ("/", "Primary positioning, trust indicators, services overview, workflow, benefits, performance narrative, testimonials, consultation CTA, and WhatsApp support."),
    ("/about", "Company overview, healthcare RCM expertise, security approach, and provider-focused value."),
    ("/services", "Medical Billing, Claims Management, Denial Management, Eligibility Verification, Payment Posting, and Revenue Analytics."),
    ("/solutions", "Provider workflow from verification through claims and payment recovery, benefits, and consultation CTA."),
    ("/resources", "RCM education, safe information-sharing guidance, and frequently asked questions."),
    ("/contact", "Support email, phone contact, and consultation request path."),
    ("/consultation", "Structured lead capture protected by Google reCAPTCHA."),
    ("/privacy, /terms, /hipaa", "Legal and trust information covering website use, privacy handling, and HIPAA-focused workflow expectations.")
]
for route, purpose in routes:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(f"{route}  "), 10.5, BLUE, bold=True)
    font(p.add_run(purpose), 10.5, NAVY)

doc.add_heading("Conversion paths", level=2)
for item in [
    "Free RCM Audit and consultation buttons route visitors to the consultation form.",
    "Navbar and footer links connect the complete public information architecture.",
    "The floating WhatsApp button and marketing-page WhatsApp card open a pre-filled conversation using a configurable number.",
    "Contact cards expose email and phone paths while warning visitors not to send protected health information through general channels.",
    "Public forms are positioned for provider/business information rather than patient care or clinical communication."
]: add_bullet(item)

doc.add_heading("Search visibility", level=2)
doc.add_paragraph("The public pages contain unique US-market titles and descriptions, canonical URLs, natural keyword targeting, Open Graph and Twitter metadata, Organization/WebSite/Service structured data, a public-only XML sitemap, and crawler controls that exclude API, CRM, authentication, and test routes. A reusable specialty SEO registry anticipates future cardiology, orthopedic, behavioral health, and primary care service pages without publishing thin placeholders.")

page_break()
doc.add_heading("2. Consultation and lead intake", level=1)
doc.add_paragraph("The consultation experience collects structured provider and practice information, calculates an internal opportunity score, applies lead priority logic, and submits through a server route. The intake design is intended to give the sales team more useful context than a basic contact form.")

doc.add_heading("Key intake capabilities", level=2)
for item in [
    "Provider/contact details, organization, specialty, NPI, billing context, claims volume, systems, challenges, and preferred contact information.",
    "Client- and server-side validation for data quality and predictable API behavior.",
    "Google reCAPTCHA v3 to reduce automated abuse.",
    "Lead scoring and priority classification to help triage follow-up.",
    "Duplicate detection to reduce repeated records based on normalized identifying fields.",
    "Confirmation and operational messaging through configured email services.",
    "Server-controlled persistence so public clients do not write directly to protected CRM collections."
]: add_bullet(item)

doc.add_heading("WhatsApp journey", level=2)
doc.add_paragraph("Visitors can start a pre-filled WhatsApp conversation from a branded floating button or marketing support card. The configured public number is normalized before the wa.me link is produced. Separately, the webhook route supports verification and inbound message processing, maintains conversation state, can extract lead details, and can create CRM leads when sufficient information is available.")
add_callout("Important distinction", "A click-to-chat button is a visitor convenience. The webhook/automation path is an operational integration and depends on correctly configured Meta credentials, webhook settings, messaging policies, and production monitoring.")

doc.add_heading("Lead lifecycle", level=2)
for step in [
    "Visitor submits a consultation or provides qualifying information through the WhatsApp workflow.",
    "Server validation, duplicate checks, scoring, and record creation occur.",
    "An authorized manager or administrator assigns the lead to an eligible salesperson.",
    "The salesperson reviews the lead, updates notes/status, and records follow-up activity.",
    "Notifications, audit activity, pipeline totals, and performance analytics reflect the operational changes."
]: add_number(step)

page_break()
doc.add_heading("3. CRM and operational capabilities", level=1)
doc.add_paragraph("The CRM is an authenticated workspace for lead and workforce operations. It includes dashboard summaries, lead lists and detail views, assignment controls, notes and follow-up, activity timelines, notifications, employee performance, user administration, role administration, and settings areas.")

doc.add_heading("Lead management", level=2)
for item in [
    "Create and edit leads through validated CRM forms.",
    "View lead cards, tabular lists, pipeline status, scoring, priority, and ownership.",
    "Assign or reassign leads only to eligible active users within the permitted scope.",
    "Capture notes, follow-up information, status changes, and activity history.",
    "Prevent unauthorized access through both API checks and Firestore rules.",
    "Notify employees about new assignments and relevant operational changes."
]: add_bullet(item)

doc.add_heading("Management and analytics", level=2)
for item in [
    "Dashboard pipeline overview and operational statistics.",
    "Aggregate totals such as assigned, unassigned, converted, and conversion rate.",
    "Employee-level assigned/completed/pending lead counts and recent activity.",
    "Role-aware visibility so users see all, team, or owned records according to permissions.",
    "User lifecycle actions including invitation/creation, updates, disablement, temporary-password reset, and performance views.",
    "Role creation and updates with validation and protection for system roles."
]: add_bullet(item)

doc.add_heading("Auditability and notifications", level=2)
doc.add_paragraph("Important lead, role, and user actions create activity records with actor, entity, action, metadata, and timestamps. Notifications are stored for user-facing alerts, can be marked read, and support assignment email coordination. This creates a foundation for accountability, operational reporting, and incident investigation, although retention and formal audit policy still require leadership approval.")

page_break()
doc.add_heading("4. Authentication, authorization, and security", level=1)
doc.add_paragraph("Firebase Authentication handles sign-in and password reset. Firebase Admin is used for privileged server operations. The application adds its own CRM user profile, role, permissions, account status, temporary-password requirements, and authorization decisions rather than treating a successful Firebase login as sufficient access.")

doc.add_heading("Authentication experience", level=2)
for item in [
    "Email/password login with neutral invalid-credential messaging.",
    "Password visibility control on the login form.",
    "Forgot-password flow using Firebase password-reset email with a neutral response that does not disclose whether an email exists.",
    "Temporary-password and mandatory personal-password-change workflow for provisioned employees.",
    "Disabled, expired, or unavailable CRM accounts are signed out and blocked from the dashboard."
]: add_bullet(item)

doc.add_heading("Authorization layers", level=2)
for item in [
    "Server API authentication verifies Firebase ID tokens and loads the active CRM user context.",
    "Permission checks control leads, analytics, activities, roles, users, and management actions.",
    "Lead visibility logic restricts access to all records, a manager's team, or a salesperson's owned leads.",
    "Firestore rules allow limited authenticated reads but deny client writes to protected CRM collections.",
    "Sensitive write paths are server-controlled through Firebase Admin and service/repository layers.",
    "Rate limiting and validation utilities exist for public or sensitive API workflows."
]: add_bullet(item)

doc.add_heading("Healthcare security and HIPAA posture", level=2)
doc.add_paragraph("The project includes HIPAA-focused messaging, least-privilege concepts, restricted public-form guidance, role-aware access, server-controlled writes, and audit-oriented activity records. These are meaningful safeguards, but they do not by themselves establish HIPAA compliance.")
add_callout("Compliance boundary", "HIPAA readiness depends on the actual services performed, data flows, workforce practices, infrastructure configuration, vendor agreements/BAAs, risk analysis, training, incident response, retention, access reviews, and evidence. Leadership should describe the product as HIPAA-focused or HIPAA-aligned unless formal legal/compliance review supports a stronger claim.", "FFF7ED")

page_break()
doc.add_heading("5. Technical architecture", level=1)
doc.add_paragraph("The application follows the Next.js App Router model. Public pages and authenticated dashboards are React routes, while server route handlers expose application APIs. Business logic is separated into validation, services, repositories, and shared types, which helps keep permissions and data access decisions outside presentation components.")

doc.add_heading("Primary technology stack", level=2)
stack = [
    ("Application", "Next.js 16.3, React 19.2, TypeScript 5"),
    ("UI", "Tailwind CSS 4, Framer Motion, Lucide React, Three.js/React Three Fiber assets"),
    ("Authentication", "Firebase Authentication and Firebase Admin"),
    ("Database", "Cloud Firestore with security rules and documented index needs"),
    ("Public protection", "Google reCAPTCHA v3 and validation/rate-limit utilities"),
    ("Messaging", "WhatsApp click-to-chat and webhook integration; Resend-based email service"),
    ("AI/data dependencies", "Google Generative AI packages and Supabase client are present for AI/search-related capabilities"),
    ("Quality gates", "ESLint, TypeScript no-emit typecheck, and Next.js production build")
]
for label, value in stack:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(f"{label}: "), 10.5, BLUE, bold=True)
    font(p.add_run(value), 10.5, NAVY)

doc.add_heading("Application layers", level=2)
for item in [
    "Pages/components: public content, consultation experience, authentication screens, and CRM views.",
    "Route handlers: HTTP validation, authentication, authorization, and response handling.",
    "Services: business rules for users, roles, activities, analytics, notifications, ownership, and visibility.",
    "Repositories: Firestore reads, transactions, batched writes, and mapping to application models.",
    "Shared libraries/types: Firebase clients, API helpers, validation, scoring, duplicate detection, ownership, permissions, SEO, and domain types."
]: add_bullet(item)

doc.add_heading("Data domains", level=2)
for item in [
    "CRM leads and consultation-derived lead data.",
    "CRM user profiles, roles, permissions, status, team/manager relationships, and password workflow state.",
    "Employee activities and audit metadata.",
    "User notifications and assignment-email coordination state.",
    "WhatsApp session/conversation state used by the webhook workflow."
]: add_bullet(item)

page_break()
doc.add_heading("6. Deployment, configuration, and operations", level=1)
doc.add_paragraph("The repository includes standard Next.js development, lint, typecheck, build, start, Firebase rules, and administrative/migration scripts. Environment variables configure Firebase, reCAPTCHA, email, WhatsApp, application URLs, and other integration credentials. Secrets must remain outside source control and should be managed through the deployment platform's encrypted environment configuration.")

doc.add_heading("Deployment checklist", level=2)
for item in [
    "Confirm the final production domain and set canonical/site URL configuration.",
    "Configure Firebase client and Admin credentials for the correct production project.",
    "Deploy and verify Firestore rules and required indexes.",
    "Configure reCAPTCHA allowed domains and production site key/secret behavior.",
    "Configure WhatsApp Business webhook verification, tokens, app secret, phone identifiers, and public click-to-chat number.",
    "Configure the production email domain/sender and validate deliverability.",
    "Run lint, typecheck, build, smoke tests, and permission-role acceptance tests.",
    "Set up monitoring for API failures, webhook failures, email failures, authentication anomalies, and Firestore usage.",
    "Document backup/export, recovery, retention, incident response, and access-review procedures."
]: add_bullet(item)

doc.add_heading("Known operational cautions", level=2)
for item in [
    "The repository README remains the default starter text and should be replaced with project-specific setup and operations documentation.",
    "Public contact details and claims should be verified before launch; placeholder phone numbers or unsupported performance statements can reduce trust.",
    "A successful build confirms compile-time integrity, not end-to-end production behavior or security/compliance certification.",
    "Webhook, email, AI, and third-party integrations require production credentials, policy compliance, monitoring, and failure-handling drills.",
    "Formal automated test coverage is not represented in the package scripts and should be added for high-risk workflows."
]: add_bullet(item)

page_break()
doc.add_heading("7. Recommended roadmap", level=1)
doc.add_heading("Priority 1 - Production readiness", level=2)
for item in [
    "Replace placeholder contact and performance claims with approved, evidence-backed information.",
    "Complete security, privacy, and HIPAA risk reviews; confirm BAAs and vendor configurations where applicable.",
    "Add end-to-end tests for consultation intake, login/reset, role permissions, lead visibility, assignment, and webhook handling.",
    "Implement centralized observability, alerting, audit retention, backup/export, and incident runbooks.",
    "Finalize the production domain, branded email delivery, analytics consent approach, and Search Console setup."
]: add_bullet(item)

doc.add_heading("Priority 2 - Operational maturity", level=2)
for item in [
    "Define lead SLAs, pipeline stage definitions, assignment rules, escalation ownership, and reporting cadence.",
    "Add controlled exports and management reports with privacy-aware access and audit trails.",
    "Schedule recurring access reviews and inactive-account cleanup.",
    "Document user onboarding/offboarding, temporary-password handling, and role-change approval.",
    "Create a staging environment with representative non-PHI test data."
]: add_bullet(item)

doc.add_heading("Priority 3 - Growth", level=2)
for item in [
    "Publish substantive specialty pages for cardiology, orthopedics, behavioral health, and primary care.",
    "Expand educational resources into authoritative guides, case studies, and measurable proof points.",
    "Connect approved scheduling, call tracking, or marketing attribution tools when business ownership is clear.",
    "Evaluate AI-assisted lead summarization or reply drafting with human review, strict data boundaries, and vendor risk approval.",
    "Develop client-facing operational reporting only after data quality, permissions, and contractual scope are mature."
]: add_bullet(item)

page_break()
doc.add_heading("8. CEO and leadership meeting Q&A", level=1)
doc.add_paragraph("The answers below are written as concise meeting responses. They should be adapted to final commercial terms, production evidence, legal advice, and approved company claims.")

doc.add_heading("Business and positioning", level=2)
qas = [
    ("What exactly is WiseMedBilling?", "WiseMedBilling is a provider-focused healthcare revenue-cycle platform and service presence that combines public lead generation with an internal CRM for managing consultations, assignments, follow-up, and performance."),
    ("Who is the target customer?", "US doctors, medical practices, clinics, healthcare providers, and healthcare organizations that need stronger billing, claims, denial, eligibility, payment-posting, or revenue visibility workflows."),
    ("What business problem are we solving?", "We reduce operational fragmentation around medical billing and lead follow-up. The public site explains the value proposition, while the CRM gives the team a controlled way to convert inquiries into managed opportunities."),
    ("What makes us different?", "The current positioning combines provider-focused RCM expertise, a security-conscious operating model, connected lead intake, and internal workflow visibility. The strongest differentiation will ultimately come from verified outcomes, specialty expertise, service quality, and client proof."),
    ("Is this a software product or a services company?", "The current project supports both a service-company market presence and an internal operations platform. Leadership should define the commercial model clearly: managed services, software-enabled services, platform licensing, or a combination."),
    ("How does the website generate revenue?", "It converts visitors through consultation requests and WhatsApp conversations, captures structured business context, and routes qualified opportunities into the CRM for follow-up."),
    ("What metrics should leadership track?", "Qualified inquiries, form completion rate, source attribution, speed to first response, assignment time, stage conversion, denial/claim outcomes for clients, sales conversion, and cost per qualified opportunity."),
]
for q, a in qas: add_qa(q, a)

doc.add_heading("Product and customer experience", level=2)
qas = [
    ("What can a visitor do today?", "Learn about the company and services, review solutions and resources, understand privacy/HIPAA expectations, contact support, request a consultation, or start a WhatsApp conversation."),
    ("What information does the consultation form collect?", "Provider and practice details, operational context, claims volume, systems, challenges, and contact preferences. Public instructions tell visitors not to submit patient-identifiable information through unapproved channels."),
    ("How are leads prioritized?", "The application calculates an internal lead score and priority using submitted business signals. Leadership should periodically validate that the scoring model matches actual conversion and client value."),
    ("How are duplicates handled?", "Normalized identifying fields are checked before creating or updating records, reducing duplicate leads and helping preserve a cleaner pipeline."),
    ("What happens after a lead is submitted?", "The server validates and stores the lead, then authorized CRM users review, assign, follow up, update status/notes, and track activity and performance."),
    ("Can the system support specialty-specific growth?", "Yes. The SEO configuration anticipates cardiology, orthopedic, behavioral health, and primary care pages. Each should be launched only with unique, credible content and operational capability."),
]
for q, a in qas: add_qa(q, a)

doc.add_heading("CRM and team operations", level=2)
qas = [
    ("Who can see which leads?", "Access is role- and ownership-aware. Administrators can have broad visibility, managers are limited to permitted team scope, and sales users are restricted to owned leads according to their permissions."),
    ("Can a salesperson see another salesperson's leads?", "Not by default. The API and Firestore rules enforce ownership and permission boundaries; any broader access must be explicitly granted through role design."),
    ("How are leads assigned?", "Authorized users select eligible active sales users. Assignment occurs in a server-side transaction, records ownership information, creates activity/notification records, and can trigger assignment email."),
    ("Can management measure employee performance?", "The CRM calculates assigned, completed, pending, activity, and conversion metrics within the viewer's permitted scope. Leadership should define agreed formulas and reporting periods before using them for compensation."),
    ("Is there an audit trail?", "Important lead, user, and role changes create activity records with actor, action, entity, metadata, and timestamps. A formal retention and review policy should still be approved."),
    ("What happens when an employee leaves?", "Accounts can be disabled and lead ownership can be managed. A documented offboarding checklist should require immediate access removal, lead reassignment, token/session review, and audit confirmation."),
]
for q, a in qas: add_qa(q, a)

page_break()
doc.add_heading("Security, privacy, and compliance", level=2)
qas = [
    ("Is the platform HIPAA compliant?", "The codebase includes HIPAA-focused safeguards and messaging, but compliance cannot be concluded from code alone. A formal risk analysis, vendor/BAA review, infrastructure configuration review, policies, training, monitoring, and evidence are required."),
    ("Does the public website collect PHI?", "It is explicitly designed for provider and business information, not patient data. Visitors are instructed not to submit patient-identifiable information unless a secure approved workflow and required agreements are in place."),
    ("How is access controlled?", "Firebase verifies identity; the application verifies active CRM status, role, permissions, password-change state, and lead scope; Firestore rules add a further restriction layer."),
    ("Are passwords stored in Firestore?", "No. Firebase Authentication manages credentials. CRM records store account workflow attributes, not plaintext passwords."),
    ("Does password reset reveal whether a user exists?", "The forgot-password UI always displays the same neutral confirmation for a valid submission, regardless of Firebase's response."),
    ("How are secrets protected?", "Secrets are expected in environment configuration and excluded from source control. Production secrets should be managed in the deployment platform, rotated, access-controlled, and never copied into documents or tickets."),
    ("What are the main security gaps before launch?", "Independent security testing, production configuration review, monitoring/alerting, incident response, backup/recovery evidence, vendor agreements, access reviews, and formal HIPAA/privacy assessment."),
]
for q, a in qas: add_qa(q, a)

doc.add_heading("Technology and scalability", level=2)
qas = [
    ("Why Next.js and Firebase?", "They provide a productive full-stack model: server-rendered/public pages, route handlers, managed authentication, and scalable document storage. The tradeoff is the need for careful authorization, query/index planning, and cost monitoring."),
    ("Can this scale?", "The architecture can scale beyond an initial team, but scale depends on Firestore query design, indexes, pagination, background processing, webhook throughput, monitoring, and cost controls."),
    ("Is the code modular?", "Yes. UI, route handlers, services, repositories, shared libraries, and types are separated. This supports incremental testing and replacement of individual integrations."),
    ("What third parties are involved?", "Firebase/Google services, reCAPTCHA, WhatsApp/Meta, email delivery, and AI/data client dependencies are present. Each production use requires credential control, vendor review, and appropriate agreements."),
    ("How do we know releases are safe?", "Current automated gates are lint, TypeScript typecheck, and production build. These catch code-quality and compile issues but should be supplemented with unit, integration, end-to-end, security, and permission regression tests."),
    ("What happens if WhatsApp or email fails?", "The application includes failure paths and notification coordination, but production reliability requires alerts, retries where safe, dashboards, and operational ownership for unresolved failures."),
]
for q, a in qas: add_qa(q, a)

doc.add_heading("SEO, growth, and credibility", level=2)
qas = [
    ("Is the site ready for search engines?", "The technical foundation is in place: unique metadata, canonicals, structured data, sitemap, robots controls, headings, and descriptive image text. Ranking still requires authority, quality content, links, performance, and time."),
    ("Which keywords are targeted?", "Healthcare revenue cycle management, medical billing services/company, RCM services, claims and denial management, eligibility verification, payment posting, healthcare analytics, outsourcing, and provider-focused high-intent terms."),
    ("Will adding many specialty pages improve rankings?", "Only if each page provides substantial and credible specialty-specific value. Thin or repetitive pages can weaken quality signals rather than improve them."),
    ("What trust assets are still needed?", "Verified company address/contact data, leadership/team profiles, client-approved case studies, measured outcomes, certifications or attestations where valid, strong privacy/security documentation, and consistent brand identity."),
    ("What should we avoid saying publicly?", "Avoid guarantees, unsupported accuracy/revenue statistics, unqualified HIPAA-compliant claims, fabricated testimonials, or service claims that operations cannot substantiate."),
]
for q, a in qas: add_qa(q, a)

page_break()
doc.add_heading("9. Meeting preparation checklist", level=1)
doc.add_heading("Decisions leadership should make", level=2)
for item in [
    "Confirm the commercial model and priority customer segments.",
    "Approve final service scope, claims, proof points, and public contact information.",
    "Name owners for sales operations, CRM administration, security, privacy/compliance, integrations, and incident response.",
    "Define lead response SLAs, pipeline stages, assignment rules, and conversion reporting.",
    "Approve the production domain, deployment environment, vendor list, and data-handling boundaries.",
    "Fund production testing, monitoring, backup/recovery, and compliance evidence work."
]: add_bullet(item)

doc.add_heading("Questions the project team should be ready to demonstrate", level=2)
for item in [
    "Can a visitor move from the homepage to a consultation and receive a neutral, usable confirmation?",
    "Can a manager assign a lead while a salesperson remains restricted to owned leads?",
    "Do activity and notification records appear after key lead actions?",
    "Does forgot-password remain neutral for existing and non-existing email addresses?",
    "Are CRM/auth/API routes absent from the sitemap and blocked in robots rules?",
    "Does WhatsApp open the correct configured number and pre-filled message without overlapping reCAPTCHA?",
    "Can the team explain which data is allowed in public channels and which requires an approved secure workflow?",
    "Can the release pass lint, typecheck, build, and agreed end-to-end acceptance tests?"
]: add_bullet(item)

doc.add_heading("Suggested 45-minute meeting flow", level=2)
for step in [
    "5 minutes - Business objective, target customer, and success definition.",
    "8 minutes - Public website and consultation journey demonstration.",
    "10 minutes - CRM lead assignment, follow-up, permissions, and analytics demonstration.",
    "8 minutes - Security, privacy, HIPAA boundary, and operational controls.",
    "6 minutes - SEO/growth foundation and specialty expansion plan.",
    "8 minutes - Decisions, owners, risks, and next milestones."
]: add_number(step)

doc.add_heading("Document scope and assurance note", level=2)
doc.add_paragraph("This briefing is based on the local project repository and completed build validations available at the time of review. It is a technical and product overview, not legal advice, a security audit, a penetration test, a HIPAA certification, a production uptime attestation, or evidence of commercial outcomes. Statements about compliance, performance, and customer results should be approved and supported before external use.")

doc.core_properties.title = "WiseMedBilling - Complete Project Briefing and CEO/Team Meeting Q&A"
doc.core_properties.subject = "Executive overview of the WiseMedBilling public website, CRM, integrations, security, SEO, and meeting preparation"
doc.core_properties.author = "WiseMedBilling"
doc.core_properties.keywords = "WiseMedBilling, healthcare RCM, medical billing, CRM, project briefing"

doc.save(OUT)
print(OUT.resolve())
